
import re
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt"}

INVALID_NAME_PHRASES = {
    "professional summary", "summary", "career objective", "objective",
    "work experience", "professional experience", "employment history",
    "education", "technical skills", "skills", "skills summary", "projects",
    "personal projects", "certifications", "certificates", "achievements",
    "contact", "contact information", "profile", "profile summary",
    "software engineer", "software developer", "full stack developer",
    "backend developer", "frontend developer", "data scientist",
    "machine learning engineer", "artificial intelligence", "resume",
    "curriculum vitae",
}

INVALID_NAME_WORDS = {
    "engineer", "developer", "intern", "student", "manager", "architect",
    "consultant", "analyst", "professional", "summary", "experience",
    "education", "skills", "projects", "certification", "objective",
    "university", "college", "email", "phone", "linkedin", "github",
    "address", "bengaluru", "bangalore", "chennai", "hyderabad", "mumbai",
    "delhi", "pune", "india",
}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if not text.strip():
        raise ValueError(
            "No extractable text found; this may be a scanned/image-only PDF."
        )

    return text


def _unique(items):
    result, seen = [], set()
    for item in items:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def _is_plausible_name(line: str) -> bool:
    clean = re.sub(r"\s+", " ", line).strip(" |•\t-")

    if not clean or len(clean) > 40:
        return False

    lower = clean.lower()

    if lower in INVALID_NAME_PHRASES:
        return False

    if any(term in lower for term in (
        "http", "www.", "@", "linkedin", "github.com",
        "phone:", "email:", "address:", "resume", "curriculum vitae",
    )):
        return False

    if re.search(r"\d|[/\\|:;@_]", clean):
        return False

    words = clean.split()
    if not 2 <= len(words) <= 4:
        return False

    if any(word.lower().strip(".,") in INVALID_NAME_WORDS for word in words):
        return False

    for word in words:
        part = word.strip(".-'")
        if not part or not re.fullmatch(r"[A-Za-z]+", part):
            return False
        if len(part) == 1 and not word.endswith("."):
            return False

    return True


def parse_candidate(path: Path, text: str) -> dict:
    emails = re.findall(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I
    )

    github_urls = re.findall(
        r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9-]+)", text, re.I
    )
    github_urls = [
        username for username in github_urls
        if username.lower() not in {"features", "topics", "settings", "orgs"}
    ]

    name = path.stem.replace("_", " ").replace("-", " ").strip()
    lines = [
        re.sub(r"\s+", " ", line).strip(" |•\t")
        for line in text.splitlines()
    ]
    lines = [line for line in lines if line]

    for line in lines[:20]:
        if _is_plausible_name(line):
            name = re.sub(r"\s+", " ", line).strip()
            break

    skill_patterns = [
        "python", "fastapi", "django", "flask", "asyncio",
        "postgresql", "postgres", "redis", "docker", "gcp",
        "google cloud", "aws", "azure", "react", "next.js",
        "javascript", "typescript", "java", "spring boot",
        "langchain", "langgraph", "llamaindex", "llama index",
        "rag", "retrieval augmented generation", "embeddings",
        "vector search", "chromadb", "pinecone", "faiss",
        "openai", "llm", "agentic", "tool calling", "multi-agent",
        "evaluation pipeline", "pytest", "unit testing", "caching",
        "observability", "concurrency", "kafka", "celery",
    ]

    lower = text.lower()
    matched = [skill for skill in skill_patterns if skill in lower]

    return {
        "candidate_name": name,
        "email": emails[0] if emails else None,
        "github_username": github_urls[0] if github_urls else None,
        "matched_skills": matched,
        "raw_text": text,
        "source_file": path.name,
    }


def list_resume_files(directory: Path):
    return sorted(
        [
            path for path in directory.iterdir()
            if path.is_file() and path.suffix.lower() in SUPPORTED
        ],
        key=lambda path: path.name.lower(),
    )
